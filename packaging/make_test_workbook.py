"""A test workbook for the built app's smoke test: 400 sales rows with 3 planted problems at known sheet rows
(12: Units -7, 32: Region "East ", 42: Total that doesn't add up). Usage: python make_test_workbook.py out.xlsx"""
import datetime
import random
import sys

import openpyxl

r = random.Random(7)
wb = openpyxl.Workbook()
ws = wb.active
ws.append(["Order ID", "Order Date", "Region", "Units", "Unit Price", "Subtotal", "Tax", "Total"])
for k in range(400):
    u, p = r.randint(1, 20), r.choice([3.6, 7.25, 9.75, 12.99, 18.0, 24.5])
    sub = round(u * p, 2)
    tax = round(sub * 0.08, 2)
    ws.append([100001 + k, datetime.date(2024, 1, 1) + datetime.timedelta(days=r.randint(0, 364)),
               r.choice(["East", "West", "North", "South"]), u, p, sub, tax, round(sub + tax, 2)])
ws["D12"], ws["C32"], ws["H42"] = -7, "East ", 999.99
wb.save(sys.argv[1])
