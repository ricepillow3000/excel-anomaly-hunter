"""report.xlsx: Data (every row, colored), Anomalies (flagged, worst first), Summary."""
from collections import Counter

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from anomaly_hunter.engine.detectors import K

FILL = {s: PatternFill("solid", fgColor=c) for s, c in
        {"Noted": "FFFFF2CC", "Low": "FFFFF2A8", "Medium": "FFFFCC80", "High": "FFFF8A80"}.items()}
RANK = {"High": 3, "Medium": 2, "Low": 1}


def write_report(df, rows, status, out_path, limits_path, warnings):
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    head = [*df.columns, "Severity", "Bucket", "Reason"]
    data.append(head)
    vals = df.astype(object).where(df.notna(), None).values.tolist()
    for v, r in zip(vals, rows):
        data.append(v + [r["severity"] or "", r["bucket"] or "", r["reason"]])
        if r["severity"]:
            for cell in data[data.max_row]:
                cell.fill = FILL[r["severity"]]
            # ponytail: one note on Reason cell, not one per triggering cell
            data.cell(data.max_row, len(head)).comment = Comment(r["reason"], "Anomaly Hunter")

    bad = wb.create_sheet("Anomalies")
    bad.append(head + ["Row"])
    order = sorted((i for i, r in enumerate(rows) if r["severity"] in RANK),
                   key=lambda i: (-RANK[rows[i]["severity"]], -rows[i]["magnitude"]))
    for i in order:
        bad.append(vals[i] + [rows[i]["severity"], rows[i]["bucket"], rows[i]["reason"], i + 2])

    sev, bkt = Counter(r["severity"] for r in rows), Counter(r["bucket"] for r in rows)
    lines = [["Counts by severity"], *([k, sev[k]] for k in ("High", "Medium", "Low", "Noted")),
             [], ["Counts by bucket"], *([k, bkt[k]] for k in ("Duplicates", "Irregularities", "Behavioral")),
             [], ["Detectors"], *status.items(),
             [], ["Limits file", str(limits_path)], ["K", K]]
    if warnings:
        lines += [[], ["Warnings"], *([w] for w in warnings)]
    summary = wb.create_sheet("Summary")
    for line in lines:
        summary.append(line)
    for ws in wb:  # a scanned cell "=HYPERLINK(...)" or "=cmd|..." stays text: opening the report never runs it
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    wb.save(out_path)
