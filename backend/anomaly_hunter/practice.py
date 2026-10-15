"""The "Start here" workbook the installer opens (instead of a blank Excel and a message that's gone after one click):
the 5 clicks that add Anomaly Hunter to Excel stay on screen while the user does them, and a Practice sheet with 5
planted problems shows what a scan finds before they try it on their own data."""
import datetime

import openpyxl
from openpyxl.styles import Font

STEPS = [
    ("Anomaly Hunter - start here", 20),
    ("", 0),
    ("First time only: add it to Excel (5 clicks)", 14),
    ("1.  Click the Home tab (top left of Excel).", 0),
    ("2.  Click Add-ins (at the right end of the Home tab).", 0),
    ("3.  Click More Add-ins.", 0),
    ("4.  Click SHARED FOLDER (at the top of the window that opens).", 0),
    ("5.  Click Anomaly Hunter, then click Add.", 0),
    ("     No SHARED FOLDER? Close every Excel window, open Excel again, and repeat from step 1.", 0),
    ("", 0),
    ("Try it (1 minute)", 14),
    ("6.  Click the Practice tab at the bottom of this window.", 0),
    ("7.  Click the Anomaly Hunter button on the Home tab. A panel opens on the right.", 0),
    ("8.  Click Find problems. It finds the 5 problems hidden in the Practice sheet:", 0),
    ("     a minus sign, a number far too big, a row typed twice, 'east ' typed differently, and an empty cell.", 0),
    ("9.  Click a problem to see why and a suggested fix. Nothing changes until you click Accept.", 0),
    ("", 0),
    ("Your own sheets", 14),
    ("Open any workbook with a row of column names on top, click Anomaly Hunter, then Find problems.", 0),
    ("Your data stays on this PC.", 0),
]
PLANTED = {11: "Units -5 (a minus sign)", 22: "Units 7777 (far too big)", 51: "a copy of row 50",
           81: 'Region "east " (typed differently)', 121: "Units left empty"}  # sheet rows


def make(path):
    """Write the workbook: "Start here" (active, read first) and "Practice" (300 orders, PLANTED problems)."""
    wb = openpyxl.Workbook()
    start = wb.active
    start.title = "Start here"
    start.column_dimensions["A"].width = 110
    for r, (text, size) in enumerate(STEPS, 1):
        start.cell(r, 1, text).font = Font(size=size or 12, bold=bool(size))
    practice = wb.create_sheet("Practice")
    practice.append(["OrderDate", "Region", "Units", "UnitCost", "Total"])
    for k in range(300):
        units, cost = 40 + (k * 7) % 23, 1.99 + k % 4
        practice.append([datetime.date(2026, 1, 6) + datetime.timedelta(days=k * 2 // 5), ["East", "West", "Central"][k % 3],
                         units, cost, round(units * cost, 2)])
    practice["C11"], practice["C22"], practice["B81"], practice["C121"] = -5, 7777, "east ", None
    for c, cell in enumerate(practice[50], 1):
        practice.cell(51, c, cell.value)
    for col, w in zip("ABCDE", (12, 10, 8, 10, 10)):
        practice.column_dimensions[col].width = w
    for row in practice.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    wb.save(path)
    return path
