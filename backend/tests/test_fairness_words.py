"""AH-18: the fairness report says only what its numbers show. Each case is a synthetic sheet from the AH-17 bias hunt
(known truth, planted by seed): too few rows is never "even", a column is an outcome only when it is one, the public
context matches the outcome and its direction, the group that stands out is named against the rest, a doubled small
rate is a real gap, small groups left out are named, and a race time or an order state is not a person."""
import random

import numpy as np
import pandas as pd
import pytest
from scipy.stats import chi2_contingency

from anomaly_hunter.agents import scan_agent
from anomaly_hunter.engine import fairness
from anomaly_hunter.engine.load import load_from_records


def people(n, rates, seed, col="Sex", outcome="Approved", groups=("F", "M"), label=None):
    """n people, group = groups[i % len], outcome 0/1 with rates[group] (label: (yes, no) text instead of 1/0)."""
    rng = np.random.default_rng(seed)
    rows = []
    for i in range(n):
        g = groups[i % len(groups)]
        hit = rng.random() < rates[g]
        rows.append([f"P{i:05d}", g, round(float(rng.normal(500, 80))), (label[0] if hit else label[1]) if label else int(hit)])
    return ["Person ID", col, "Score", outcome], rows


def report(cols, rows):
    return scan_agent.scan(cols, rows)[-1]


def test_a_big_spread_chance_explains_is_no_clear_difference_never_even():
    # AH-17 W1: 80 people, 15% vs 40% planted (20% vs 42% drawn, raw p 0.054) read "even across groups (... within
    # normal variation)". Chance explains it (raw p >= 0.05), so it is not "unclear" either (Council A, AH-18 review)
    cols, rows = people(80, {"F": 0.15, "M": 0.4}, 80)
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Approved rate")
    assert t["verdict"] == "even" and t["p"] >= 0.05 and "even" not in t["text"] and "no clear difference" in t["text"], t["text"]
    drawn = sorted(np.mean([r[3] for r in rows if r[1] == g]) for g in "FM")
    # raw p 0.054 is near the line: "within normal variation" claimed more than the numbers show (Council A, AH-20)
    approved = sum(r[3] for r in rows)  # 25 of 80: the approvals are what there is little of
    assert t["p"] < 0.10 and t["text"].endswith(f"no clear difference across groups ({drawn[0]:.0%} to {drawn[1]:.0%}); "
                                                 f"with only {approved} Approved = 1 to go on, a gap this size could be chance"), t["text"]
    assert "within normal variation" not in t["text"]
    # a difference that is small as well as unproven is plain "even"
    t = next(t for t in report(*people(500, {"F": 0.3, "M": 0.3}, 7))["tests"] if t["test"] == "Approved rate")
    assert t["verdict"] == "even" and "even across groups" in t["text"] and "within normal variation" in t["text"]


@pytest.mark.parametrize("n, rates, seed", [(3000, (0.010, 0.018), 2), (2000, (0.02, 0.032), 29), (6000, (0.005, 0.009), 0)])
def test_near_the_line_on_a_big_sheet_the_few_are_the_events_never_the_rows(n, rates, seed):
    # 3,000 rows read "with these few rows a gap this size could be chance": what is few is the rows flagged or the
    # defaults, not the rows (Council B, AH-20 review). The sentence counts what there is little of, and counts it right
    rng = np.random.default_rng(seed)
    rows = [[f"P{i}", "F" if i % 2 else "M", round(float(rng.normal(50000, 8000))), int(rng.random() < (rates[0] if i % 2 else rates[1]))]
            for i in range(n)]
    tests = report(["ID", "Sex", "Income", "Defaulted"], rows)["tests"]
    near = [t for t in tests if "could be chance" in t["text"]]
    assert near and not any("few rows" in t["text"] for t in tests), [t["text"] for t in tests]
    for t in near:
        events = round(sum(size * rate for size, rate in t["groups"].values()))
        rarer = min(events, sum(size for size, _ in t["groups"].values()) - events)
        assert f"with only {rarer} " in t["text"] and t["text"].endswith(" to go on, a gap this size could be chance"), t["text"]


def test_a_big_spread_with_no_evidence_at_all_is_within_normal_variation():
    # a big spread whose raw p is far from any line (0.10+) is what chance gives: "within normal variation" stays
    for seed in range(200):
        cols, rows = people(40, {"F": 0.3, "M": 0.3}, 500 + seed)
        t = next((t for t in report(cols, rows)["tests"] if t["test"] == "Approved rate"), None)
        if t and t["p"] >= 0.10 and fairness._either(*sorted(v[1] for v in t["groups"].values())[::-1]):
            assert t["verdict"] == "even" and t["text"].endswith("within normal variation)"), t["text"]
            assert "no clear difference" in t["text"] and "could be chance" not in t["text"], t["text"]
            return
    raise AssertionError("no big spread with p >= 0.10 drawn")


def test_a_big_gap_the_raw_evidence_suggests_but_cannot_prove_is_not_enough_rows_to_tell():
    # raw p under 0.05 but not significant after the correction, and a big gap: "not enough rows to tell"
    cols, rows = people(80, {"F": 0.15, "M": 0.4}, 10)  # drawn 15% vs 38%, raw p 0.042
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Approved rate")
    assert 0.01 <= t["p"] < 0.05, t["p"]
    assert t["verdict"] == "unclear" and "not enough rows to tell" in t["text"] and "even" not in t["text"], t["text"]


@pytest.mark.parametrize("groups, n", [([f"C{k:02d}" for k in range(40)], 8000), (["N", "S", "E", "W"], 800), (["F", "M"], 600)])
def test_noise_is_called_even_unless_its_raw_evidence_is_suggestive(groups, n):
    # Council A: pure noise over 20 seeds read "not enough rows to tell" 13/20 (40 groups), 5/20 (4), 1/20 (2): the gap was
    # judged on the group picked for standing out, with no p threshold (p 0.74 said "not enough rows to tell")
    unclear = 0
    for seed in range(20):
        cols, rows = people(n, {g: 0.3 for g in groups}, 100 + seed, col="Country" if len(groups) > 2 else "Sex", groups=groups)
        for t in report(cols, rows)["tests"]:
            if t["test"] == "Approved rate":
                suggestive = t["p"] < 0.05 and (len(groups) == 2 or t["p_group"] * len(groups) < 0.05)
                assert t["verdict"] != "gap" and (t["verdict"] == "even" or suggestive), (seed, t["p"], t.get("p_group"), t["text"])
                unclear += t["verdict"] == "unclear"
    assert unclear <= 2, unclear


def test_every_test_says_which_kind_of_check_it_is():
    # the pane names only the checks that ran (frontend/selfcheck.js): it needs each test's kind
    cols, rows = people(500, {"F": 0.3, "M": 0.3}, 7)
    rows[3][2] = 5000.0  # a x10 slip, so rows are flagged
    for i in range(0, 60, 2):
        rows[i][2] = ""  # blanks
    out = report(cols, rows)
    assert {t["test"]: t["measure"] for t in out["tests"]} == {"flag rate": "flag", "missing Score": "missing", "Approved rate": "outcome"}


@pytest.mark.parametrize("yes, no", [("Approved", "Denied"), ("Yes", "No"), ("Pass", "Fail"), (">50K", "<=50K"), ("Hired", "Not hired"),
                                     ("True", "False"), ("Defaulted", "Repaid"), ("Declined", "Approved"), ("Defaulted", "Paid"),
                                     ("Delinquent", "Current"), ("Accepted", "Declined"), ("Fraud", "Legit"), ("Active", "Inactive"),
                                     ("Win", "Lose"), ("Left", "Stayed"), ("Graduated", "Dropped out"), ("Spam", "Ham"),
                                     # one-letter results, like Y/N and T/F (AH-20); and a contract type, which an employer decides
                                     ("P", "F"), ("W", "L"), ("Part-time", "Full-time")])
def test_a_two_label_column_of_opposites_is_an_outcome(yes, no):
    cols, rows = people(600, {"F": 0.3, "M": 0.3}, 3, outcome="Result", label=(yes, no))
    assert any(t["test"].startswith("Result ") and t["measure"] == "outcome" for t in report(cols, rows)["tests"])


@pytest.mark.parametrize("col, labels", [("Shift", ("Night", "Day")), ("Dept", ("A", "B")), ("Hand", ("Left", "Right")),
                                         ("Session", ("AM", "PM")), ("When", ("Weekday", "Weekend")), ("Setting", ("Urban", "Rural")),
                                         # AH-20: when, how and where, never a result
                                         ("Shift", ("Night", "Morning")), ("Payment", ("Card", "Cash")), ("Channel", ("Wholesale", "Retail")),
                                         ("Device", ("PC", "Mac")), ("Term", ("Fall", "Spring")), ("Term", ("Autumn", "Spring")),
                                         ("Survey", ("Post", "Pre"))])
def test_a_two_label_column_that_is_not_an_outcome_is_not_tested_as_one(col, labels):
    # AH-17 W3: "Shift" Day/Night by Gender was "a real gap, check why" with a census sentence about earning over 50K
    cols, rows = people(1000, {"F": 0.6, "M": 0.27}, 4, col="Gender", outcome=col, label=labels)
    tests = report(cols, rows)["tests"]
    assert not any(t["test"].startswith(col + " ") for t in tests), [t["text"] for t in tests]


def test_public_context_is_shown_only_for_an_income_gap_in_the_measured_direction():
    # the census sentence measures who earns over 50K (men more than women); it fits that outcome, that way round only
    cols, rows = people(1000, {"Male": 0.45, "Female": 0.25}, 10, col="Gender", outcome="Loan Status", groups=("Female", "Male"),
                        label=("Denied", "Approved"))
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Loan Status Denied rate")
    assert t["verdict"] == "gap" and "context" not in t, t.get("context")  # a loan, not earnings
    cols, rows = people(1000, {"Male": 0.4, "Female": 0.2}, 11, col="Sex", outcome="Income", groups=("Female", "Male"), label=(">50K", "<=50K"))
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Income >50K rate")
    assert t["verdict"] == "gap" and "30.4% of men" in t["context"]  # earnings, men higher: as measured
    cols, rows = people(1000, {"Male": 0.2, "Female": 0.4}, 12, col="Sex", outcome="Income", groups=("Female", "Male"), label=(">50K", "<=50K"))
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Income >50K rate")
    assert t["verdict"] == "gap" and "context" not in t, t.get("context")  # the other way round: the census says nothing about it


def test_with_many_groups_the_group_that_stands_out_is_named_against_the_rest_with_its_own_p_value():
    # AH-17: 40 countries, C07 at 60% and the rest at 30% read "Country C07: 56% vs 23% for C03" - C03 is an ordinary
    # country that happened to land lowest, and the p-value tested all 40 at once
    countries = [f"C{k:02d}" for k in range(40)]
    cols, rows = people(8000, {**{c: 0.3 for c in countries}, "C07": 0.6}, 12, col="Country", groups=countries)
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Approved rate")
    assert t["verdict"] == "gap" and t["text"].startswith("Country C07:") and "for the other groups" in t["text"], t["text"]
    assert not any(f"for C{k:02d}" in t["text"] for k in range(40))
    g = pd.Series([r[1] for r in rows]) == "C07"
    y = pd.Series([r[3] for r in rows]) == 1
    assert t["p_group"] == pytest.approx(chi2_contingency(pd.crosstab(g, y))[1], rel=1e-9)
    rest = y[~g].mean()
    assert f"vs {rest:.0%} for the other groups" in t["text"]


def test_a_doubled_small_rate_is_a_real_gap_not_a_small_difference():
    # AH-17 W2: 2% vs 4% defaults on 20,000 people (p ~ 0) read "close across groups (2% to 4% - a real difference, but small)"
    cols, rows = people(20000, {"F": 0.02, "M": 0.04}, 3, outcome="Defaulted")
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Defaulted rate")
    assert t["verdict"] == "gap" and t["text"].startswith("Sex M:"), t["text"]


def test_small_groups_left_out_are_named():
    # AH-17: 15 rows of X, all approved, were dropped with no word; the report read as if every group was compared
    rng = np.random.default_rng(8)
    rows = [[f"P{i}", "X" if i < 15 else ("F" if i % 2 else "M"), round(float(rng.normal(50000, 8000))), int(i < 15 or rng.random() < 0.3)]
            for i in range(600)]
    t = next(t for t in report(["Person ID", "Sex", "Income", "Approved"], rows)["tests"] if t["test"] == "Approved rate")
    assert "X (15 rows)" in t["text"] and "too few to compare" in t["text"], t["text"]
    # every group too small: said, not skipped, and no verdict either way
    rows = [[f"P{i}", "F" if i % 2 else "M", 50000 + i, int(i % 3 == 0)] for i in range(30)]
    t = next(t for t in report(["Person ID", "Sex", "Income", "Approved"], rows)["tests"] if t["test"] == "Approved rate")
    assert t["verdict"] == "unclear" and "F (15 rows), M (15 rows)" in t["text"], t["text"]


def test_a_race_time_or_an_order_state_is_not_read_as_a_person():
    rng = random.Random(4)
    df, types, _ = load_from_records(["Runner", "Race Time", "Club", "Finished"],
                                     [[i, round(rng.gauss(60, 8), 1), "AB"[i % 2], int(rng.random() < 0.9)] for i in range(300)])
    assert fairness.report(df, types, [{"severity": None}] * len(df))["protected"] == {}
    df, types, _ = load_from_records(["Order", "State", "Amount"], [[i, ["Pending", "Shipped", "Returned"][i % 3], 50 + i % 9] for i in range(300)])
    assert fairness.report(df, types, [{"severity": None}] * len(df))["protected"] == {}
    # a person's race and a US state still are
    df, types, _ = load_from_records(["ID", "Race", "Applicant race", "State", "Score"],
                                     [[i, "AB"[i % 2], "CD"[i % 2], ["Texas", "Ohio", "Utah"][i % 3], i % 7] for i in range(300)])
    assert set(fairness.report(df, types, [{"severity": None}] * len(df))["protected"]) == {"Race", "Applicant race", "State"}


def test_the_excel_report_says_not_enough_rows_to_tell_never_even_for_an_unclear_test(tmp_path):
    from openpyxl import load_workbook

    from anomaly_hunter.engine.report import write_report
    cols, rows = people(80, {"F": 0.15, "M": 0.4}, 10)  # raw p 0.042: "not enough rows to tell"
    small = [[f"Q{i}", "X", 500.0, 1] for i in range(15)]  # a group too small to compare: no p-value at all
    df, out, *_, fair = scan_agent.scan(cols, rows + small)
    write_report(df, out, {}, tmp_path / "r.xlsx", None, [], fair)
    lines = [[c.value for c in r] for r in load_workbook(tmp_path / "r.xlsx")["Fairness"].iter_rows()]
    unclear = [r for r in lines if r[1] == "Sex: Approved rate"]
    assert unclear and unclear[0][0] == "Not enough rows to tell", unclear


@pytest.mark.parametrize("yes, no", [("Declined", "Approved"), ("Defaulted", "Paid"), ("Delinquent", "Current")])
def test_a_real_loan_gap_is_reported_whatever_words_the_outcome_uses(yes, no):
    # Council A: a 15% vs 35% gap written Declined/Approved, Defaulted/Paid or Delinquent/Current was lost in AH-18
    cols, rows = people(2000, {"Female": 0.15, "Male": 0.35}, 31, col="Gender", outcome="Loan Status", groups=("Female", "Male"), label=(yes, no))
    t = next(t for t in report(cols, rows)["tests"] if t["test"].startswith("Loan Status "))
    assert t["verdict"] == "gap" and t["text"].startswith("Gender Male:"), t["text"]


def test_rates_under_ten_percent_keep_one_decimal():
    # Council A: 1.2% vs 0.4% read "1% ... vs 0%", and 1.0% vs 1.5% read "(1% to 2%)"
    cols, rows = people(40000, {"F": 0.004, "M": 0.012}, 3, outcome="Defaulted")
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Defaulted rate")
    drawn = {g: np.mean([r[3] for r in rows if r[1] == g]) for g in "FM"}
    assert f"{drawn['M']:.1%}" in t["text"] and f"{drawn['F']:.1%}" in t["text"], t["text"]
    cols, rows = people(40000, {"F": 0.010, "M": 0.015}, 3, outcome="Defaulted")
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Defaulted rate")
    drawn = sorted(np.mean([r[3] for r in rows if r[1] == g]) for g in "FM")
    assert f"{drawn[0]:.1%}" in t["text"] and f"{drawn[1]:.1%}" in t["text"], t["text"]


def test_the_same_people_get_the_same_verdict_however_their_outcome_is_written():
    # Council A: Approved = 1 at 89% vs 74% read "small", the same people as Approved/Denied read a gap (26% vs 11% denied)
    cols, rows = people(4000, {"F": 0.90, "M": 0.75}, 5, outcome="Approved")
    as_number = next(t for t in report(cols, rows)["tests"] if t["test"] == "Approved rate")
    words = [[*r[:3], "Approved" if r[3] else "Denied"] for r in rows]
    as_words = next(t for t in report(cols, words)["tests"] if t["test"].startswith("Approved "))
    assert as_number["verdict"] == as_words["verdict"] == "gap", (as_number["text"], as_words["text"])


def test_two_groups_that_stand_out_together_are_both_named():
    # Council A: N and S at 40%, E and W at 20% read "N: 40% vs 26% for the other groups", as if N alone stood out
    cols, rows = people(4000, {"N": 0.4, "S": 0.4, "E": 0.2, "W": 0.2}, 14, col="Region", groups=("N", "S", "E", "W"))
    t = next(t for t in report(cols, rows)["tests"] if t["test"] == "Approved rate")
    assert t["verdict"] == "gap" and t["text"].startswith("Region N and S:"), t["text"]
    assert "for the other groups" in t["text"]
