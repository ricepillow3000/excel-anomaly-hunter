"""AH-63 (blind round 38): the fairness report said what its numbers could not mean. Blanks by state that only say which
states a source tracked ("State NV: 64% missing Trump Endorsed? vs 95% - a real gap"), a later survey question blank
because an earlier answer skipped it ("Gender Female: 28% missing Which character shot first? vs 15%"), "LGBTQ?" never
read as a group, a nationality of 89 labels dropped without a word, and a birth day or a death country taken for a group
of people. Synthetic sheets only, shaped like those tables."""
import numpy as np
import pytest

from anomaly_hunter.agents.scan_agent import scan

STATES = ["CA", "TX", "NV", "OK", "TN", "MD", "NY", "PA"]


def fair(cols, rows):
    return scan(cols, rows)[5]


def gaps(f):
    return [t["text"] for t in f["tests"] if t["verdict"] == "gap"]


def candidates(seed=1):
    # a primary's candidates: who they are, and endorsements a group recorded only for the states it tracked
    rng = np.random.default_rng(seed)
    cols = ["Candidate", "State", "General Status", "Race", "Veteran?", "LGBTQ?", "Elected Official?", "STEM?", "Won Primary",
            "Trump Endorsed?", "NRA Endorsed?"]
    rows = []
    for i in range(600):
        st = STATES[i % len(STATES)]
        lgbtq = "Yes" if rng.random() < 0.15 else "No"
        rows.append([f"Candidate {i}", st, "" if st == "OK" and rng.random() < 0.85 else str(rng.choice(["On the Ballot", "Lost"])),
                     str(rng.choice(["White", "Nonwhite"])), str(rng.choice(["Yes", "No"])), lgbtq, str(rng.choice(["Yes", "No"])),
                     str(rng.choice(["Yes", "No"])), "Yes" if rng.random() < (0.15 if lgbtq == "Yes" else 0.5) else "No",
                     str(rng.choice(["Yes", "No"])) if st == "NV" and rng.random() < 0.4 else "",
                     str(rng.choice(["Yes", "No"])) if st in ("NV", "TX") and rng.random() < 0.7 else ""])
    for i in [k for k in range(600) if rows[k][1] == "TN"][:14]:  # candidates with no profile at all
        rows[i][3:8] = [""] * 5
    return cols, rows


def test_blanks_by_state_that_say_which_states_a_source_kept_are_never_a_gap():
    f = fair(*candidates())
    said = [t["text"] for t in f["tests"] if t["column"] == "State" and t["measure"] == "missing"]
    assert said == [], said[:3]


def test_lgbtq_is_a_group_of_people_and_its_gap_is_found():
    f = fair(*candidates())
    assert "LGBTQ?" in f["protected"], f["protected"]
    assert any(t["column"] == "LGBTQ?" and t["verdict"] == "gap" and "Won Primary" in t["text"] for t in f["tests"]), gaps(f)


def test_a_real_gap_in_blanks_by_a_people_column_is_still_found():
    rng = np.random.default_rng(2)
    rows = []
    for i in range(600):
        sex = ["Female", "Male"][i % 2]
        rows.append([f"P{i}", sex, "" if rng.random() < (0.4 if sex == "Female" else 0.1) else int(rng.integers(20000, 90000)),
                     str(rng.choice(["North", "South"]))])
    f = fair(["Person", "Gender", "Income", "Region"], rows)
    assert any(t["column"] == "Gender" and t["verdict"] == "gap" and "missing Income" in t["text"] for t in f["tests"]), gaps(f)


def star_wars(skip_among_seen=(0.03, 0.03), seed=3):
    rng = np.random.default_rng(seed)
    seen_q, shot = "Have you seen any of the 6 films in the Star Wars franchise?", "Which character shot first?"
    rows = []
    for i in range(1000):
        sex = ["Female", "Male"][i % 2]
        seen = "Yes" if rng.random() < (0.6 if sex == "Female" else 0.85) else "No"
        p = skip_among_seen[0] if sex == "Female" else skip_among_seen[1]
        rows.append([f"R{i}", seen, "" if seen == "No" or rng.random() < p else str(rng.choice(["Han", "Greedo", "I don't understand this question"])),
                     sex, str(rng.choice(["18-29", "30-44", "45-60", "> 60"]))])
    return ["RespondentID", seen_q, shot, "Gender", "Age"], rows, shot


def test_a_question_an_earlier_answer_skipped_is_routing_never_a_gap():
    cols, rows, shot = star_wars()
    f = fair(cols, rows)
    said = [t["text"] for t in f["tests"] if t["test"] == f"missing {shot}" and t["column"] == "Gender"]
    assert said and not any("real gap" in s for s in said), said
    assert all("counted where it was asked (Have you seen any of the 6 films in the Star Wars franchise?: Yes)" in s for s in said), said


def test_a_real_gap_among_the_respondents_who_were_asked_is_still_found():
    cols, rows, shot = star_wars(skip_among_seen=(0.3, 0.04))
    f = fair(cols, rows)
    assert any(t["test"] == f"missing {shot}" and t["column"] == "Gender" and t["verdict"] == "gap" for t in f["tests"]), gaps(f)


def people(n=900, seed=4):
    rng = np.random.default_rng(seed)
    nations = [f"Country {k:02d}" for k in range(89)]
    return [[f"Athlete {i}", str(rng.choice(["Men", "Women"])), str(rng.choice(nations)), int(rng.integers(1, 29)),
             str(rng.choice(["USA", "Cuba", "Mexico"])), "Medal" if rng.random() < 0.3 else "None"] for i in range(n)]


def test_a_nationality_of_89_labels_is_named_as_left_out_with_its_reason():
    f = fair(["Athlete", "gender", "nationality", "birthDay", "deathCountry", "Result"], people())
    said = [t["text"] for t in f["tests"] if t["column"] == "nationality"]
    assert said == ["nationality: 89 different values, too many to compare as groups (50 at most) - left out"], said
    assert f["protected"].get("nationality") == "demographic"


@pytest.mark.parametrize("col", ["birthDay", "deathCountry"])
def test_a_day_of_birth_or_a_country_of_death_is_no_group_of_people(col):
    f = fair(["Athlete", "gender", "nationality", "birthDay", "deathCountry", "Result"], people())
    assert col not in f["protected"] and not any(t["column"] == col for t in f["tests"]), f["protected"]


def test_an_option_of_a_select_all_that_apply_question_is_never_measured_as_missing():
    # Thanksgiving: "Will you watch any of the following programs? Please select all that apply. - Macy's Parade" is blank
    # when it was not chosen: "Male: 58% missing ... vs 45% for Female - a real gap" said who watches as who skipped
    rng = np.random.default_rng(5)
    col = "Will you watch any of the following programs on Thanksgiving? Please select all that apply. - Macy's Parade"
    rows = [[f"R{i}", ["Male", "Female"][i % 2], "Macy's Parade" if rng.random() < (0.42 if i % 2 else 0.55) else "",
             str(rng.choice(["Yes", "No"]))] for i in range(1000)]
    f = fair(["RespondentID", "What is your gender?", col, "Do you celebrate Thanksgiving?"], rows)
    assert not any(t["test"] == f"missing {col}" for t in f["tests"]), [t["text"] for t in f["tests"]][:3]


def test_a_survey_routed_past_by_one_answer_and_left_by_a_few_who_stopped_is_still_routing():
    # Thanksgiving: the 78 who do not celebrate skip everything (73% of the blanks), 3% of the rest stopped answering;
    # the young celebrate less, so "Age 18 - 29: 14% missing How far will you travel? vs 5.8%" read as a gap
    rng = np.random.default_rng(6)
    cols = ["RespondentID", "Do you celebrate Thanksgiving?", "How far will you travel for Thanksgiving?", "Age"]
    rows = []
    for i in range(1058):
        age = str(rng.choice(["18 - 29", "30 - 44", "45 - 59", "60+"]))
        yes = rng.random() > (0.14 if age == "18 - 29" else 0.05)
        far = "" if not yes or rng.random() < 0.03 else str(rng.choice(["Thanksgiving is local", "Few hours", "Out of town"]))
        rows.append([f"R{i}", "Yes" if yes else "No", far, age])
    f = fair(cols, rows)
    said = [t["text"] for t in f["tests"] if t["test"] == f"missing {cols[2]}" and t["column"] == "Age"]
    assert said and not any("real gap" in s for s in said) and all("counted where it was asked" in s for s in said), said


def ships(sex=False, seed=5):
    # a ship registry: the flag's 60 nationalities, its home port and a tonnage; no person on it unless a crew's sex is
    rng = np.random.default_rng(seed)
    cols = ["Vessel", "Nationality", "Port city", "Gross tonnage"] + (["sex"] if sex else [])
    rows = [[f"Vessel {i}", f"Flag {i % 60}", ["Rotterdam", "Hamburg", "Antwerp", "Le Havre"][i % 4],
             float(round(rng.lognormal(9, 1), 0)) if i % 7 else ""] + ([str(rng.choice(["F", "M"]))] if sex else [])
            for i in range(600)]
    return cols, rows


def test_a_column_too_wide_to_compare_never_makes_a_sheet_about_people():
    # AH-63 review 1: a 60-label Nationality was kept to be named as left out, and that alone made the registry a people
    # sheet whose ports got gap tests; on main it fell out and the sheet had no report
    assert fair(*ships()) == {"protected": {}, "tests": []}  # as on main: no report
    f = fair(*ships(sex=True))
    said = [t["text"] for t in f["tests"] if t["column"] == "Nationality"]
    assert said == ["Nationality: 60 different values, too many to compare as groups (50 at most) - left out"], said


def test_a_left_out_column_is_written_in_the_downloaded_report_as_left_out_never_even(tmp_path):
    # AH-63 review 1: report.py wrote every verdict it did not know as "Even", a comparison that never ran
    from openpyxl import load_workbook

    from anomaly_hunter.engine.report import write_report
    cols, rows = ships(sex=True)
    df, out, *_, f = scan(cols, rows)
    write_report(df, out, {}, tmp_path / "r.xlsx", None, [], f)
    lines = [[c.value for c in r] for r in load_workbook(tmp_path / "r.xlsx")["Fairness"].iter_rows()]
    left = [r for r in lines if r[1] == "Nationality: groups"]
    assert left and left[0][0] == "Left out - too many groups", left
    assert not any(r[0] == "Even" and r[2] and "left out" in str(r[2]) for r in lines), lines
