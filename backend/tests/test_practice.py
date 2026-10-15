"""The installer's "Start here" workbook: the steps are what the user reads first, and the Practice sheet really has
the 5 problems the steps promise - found by a scan, nothing else flagged."""
import openpyxl

from anomaly_hunter import practice
from anomaly_hunter.agents import scan_agent


def test_start_here_opens_on_the_steps_and_practice_has_exactly_the_promised_problems(tmp_path):
    wb = openpyxl.load_workbook(practice.make(tmp_path / "Start here.xlsx"))
    assert wb.sheetnames == ["Start here", "Practice"] and wb.active.title == "Start here"
    steps = [c.value for c in wb["Start here"]["A"] if c.value]
    assert any("SHARED FOLDER" in s for s in steps) and any("Find problems" in s for s in steps)
    grid = [list(r) for r in wb["Practice"].iter_rows(values_only=True)]
    out = scan_agent.scan(grid[0], [[("" if v is None else v) for v in r] for r in grid[1:]])[1]
    assert {i + 2 for i, r in enumerate(out) if r["severity"] in ("High", "Medium", "Low")} == set(practice.PLANTED)


def test_the_installer_finds_desktop_excel_or_says_none(tmp_path, monkeypatch):
    """Excel is started by its own path (App Paths), never by "open this .xlsx": a PC without Excel showed "Pick an app"
    and held the installer until someone picked one (fresh-PC CI, 2026-10-15)."""
    import sys
    import types

    from anomaly_hunter import app
    exe = tmp_path / "EXCEL.EXE"
    exe.write_bytes(b"")
    found = {"value": None}

    class Key:
        def __init__(self, root):
            self.root = root
        def __enter__(self):
            if self.root == "HKCU" or found["value"] is None:
                raise OSError("no such key")
            return self
        def __exit__(self, *a):
            return False

    fake = types.SimpleNamespace(HKEY_CURRENT_USER="HKCU", HKEY_LOCAL_MACHINE="HKLM",
                                 OpenKey=lambda root, key: Key(root), QueryValue=lambda k, name: found["value"])
    monkeypatch.setitem(sys.modules, "winreg", fake)
    assert app.excel_exe() is None                       # not registered
    found["value"] = f'"{exe}"'
    assert app.excel_exe() == exe                        # registered (quoted, as Office writes it)
    found["value"] = str(tmp_path / "gone" / "EXCEL.EXE")
    assert app.excel_exe() is None                       # registered but uninstalled
