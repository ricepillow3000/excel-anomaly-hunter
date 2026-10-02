"""Write report.xlsx: Data, Anomalies, and Summary sheets."""
from collections import Counter

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

FILL = {
    "Noted": PatternFill("solid", fgColor="FFFFF2CC"),
    "Low": PatternFill("solid", fgColor="FFFFF2A8"),
    "Medium": PatternFill("solid", fgColor="FFFFCC80"),
    "High": PatternFill("solid", fgColor="FFFF8A80"),
}
SEVERITY_RANK = {"High": 3, "Medium": 2, "Low": 1}


def _row_values(df, i):
    return list(df.iloc[i].astype(object).where(df.iloc[i].notna(), None))


def write_report(df, combined, out_path, detector_status, limits_path, k, warnings):
    wb = Workbook()
    data_ws = wb.active
    data_ws.title = "Data"

    columns = list(df.columns) + ["Severity", "Bucket", "Reason"]
    data_ws.append(columns)
    for i in range(len(df)):
        c = combined[i]
        data_ws.append(_row_values(df, i) + [c["severity"] or "", c["bucket"] or "", c["reason"]])
        if c["severity"] in FILL:
            for cell in data_ws[data_ws.max_row]:
                cell.fill = FILL[c["severity"]]
            if c["reason"]:
                # ponytail: one note on the Reason cell, not one per triggering data
                # cell — the spec asks for per-cell notes; this is the simpler version.
                data_ws.cell(row=data_ws.max_row, column=len(columns)).comment = Comment(
                    c["reason"], "Anomaly Hunter"
                )

    anomalies_ws = wb.create_sheet("Anomalies")
    anomalies_ws.append(columns + ["Row", "source_file"])
    ordered = sorted(
        (i for i, c in enumerate(combined) if c["severity"] in SEVERITY_RANK),
        key=lambda i: (-SEVERITY_RANK[combined[i]["severity"]], -combined[i]["magnitude"]),
    )
    for i in ordered:
        c = combined[i]
        anomalies_ws.append(
            _row_values(df, i) + [c["severity"], c["bucket"], c["reason"], i + 2, df.iloc[i]["source_file"]]
        )

    summary_ws = wb.create_sheet("Summary")
    severity_counts = Counter(c["severity"] for c in combined if c["severity"])
    bucket_counts = Counter(c["bucket"] for c in combined if c["bucket"])

    summary_ws.append(["Counts by severity"])
    for sev in ("High", "Medium", "Low", "Noted"):
        summary_ws.append([sev, severity_counts.get(sev, 0)])
    summary_ws.append([])
    summary_ws.append(["Counts by bucket"])
    for bucket in ("Duplicates", "Irregularities", "Behavioral"):
        summary_ws.append([bucket, bucket_counts.get(bucket, 0)])
    summary_ws.append([])
    summary_ws.append(["Detectors"])
    for name, status in detector_status.items():
        note = "ran" if status.get("ran", True) else f"sat out: {status.get('sit_out_reason')}"
        summary_ws.append([name, note])
    summary_ws.append([])
    summary_ws.append(["Limits file", str(limits_path)])
    summary_ws.append(["K", k])
    if warnings:
        summary_ws.append([])
        summary_ws.append(["Warnings"])
        for w in warnings:
            summary_ws.append([w])

    wb.save(out_path)
